"""Bounded, browser-safe Stage 10 evidence exports."""

from __future__ import annotations

from csv import reader
from datetime import datetime, timezone
from hashlib import sha256
from html import escape
from importlib.metadata import PackageNotFoundError, version
from io import BytesIO
import json
import math
from pathlib import Path
import re
from typing import Any, Mapping
from urllib.parse import parse_qsl, urlencode, urlparse, urlunparse

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .analysis import AnalysisFrames, TEMPORAL_NOTE, build_analysis_frames, event_summary, temporal_summary
from .models import QueryResult
from .report_charts import report_charts
from .visuals import EXPOSURE_NOTE, MAP_NOTE, build_map_layers


CSV_NAMES = {
    "events.csv": "events",
    "impact_observations.csv": "impacts",
    "correlation_evidence.csv": "correlations",
    "quality_summary.csv": "quality",
}
SHEET_NAMES = (
    "Query Metadata",
    "Events",
    "Impact Observations",
    "Correlation Evidence",
    "Quality Summary",
    "Definitions",
)
SENSITIVE_PARTS = ("token", "secret", "signature", "credential", "accesskey", "authorization")
PATH_RE = re.compile(r"(?i)(?<![a-z])(?:[a-z]:[\\/]|file:/+)[^\s;,]+")
BEARER_RE = re.compile(r"(?i)\bbearer\s+[^\s;,]+")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10, color="222222")


def _software_version() -> str:
    try:
        return version("guard-pdc-explorer")
    except PackageNotFoundError:
        return "0.1.0"


def public_text(value: str) -> str:
    """Redact browser/export-unsafe paths and credential values."""

    value = PATH_RE.sub("<local-path>", value)
    value = BEARER_RE.sub("Bearer <redacted>", value)
    if "://" not in value:
        return value
    parsed = urlparse(value)
    if not parsed.query:
        return value
    query = [
        (key, "<redacted>" if any(part in key.lower() for part in SENSITIVE_PARTS) else child)
        for key, child in parse_qsl(parsed.query, keep_blank_values=True)
    ]
    return urlunparse(parsed._replace(query=urlencode(query)))


def _public_value(value: Any) -> Any:
    if isinstance(value, str):
        return public_text(value)
    if isinstance(value, (dict, list, tuple, set)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
    if value is None or value is pd.NA or isinstance(value, float) and math.isnan(value):
        return None
    if hasattr(value, "item"):
        return value.item()
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def public_frame(frame: pd.DataFrame) -> pd.DataFrame:
    """Return evidence suitable for browser display and downloads."""

    public = frame.drop(
        columns=[column for column in ("endpoint_or_file", "raw_pointer", "geometry", "assets", "original_properties") if column in frame],
        errors="ignore",
    ).copy()
    for column in public.columns:
        public[column] = public[column].map(_public_value)
    return public


def _joined(values: pd.Series) -> str:
    return ", ".join(sorted({str(value) for value in values if value is not None and not pd.isna(value)}))


def _with_evidence(frame: pd.DataFrame, provenance: pd.DataFrame, entity_type: str, id_column: str) -> pd.DataFrame:
    if frame.empty or provenance.empty:
        return frame.copy()
    evidence = provenance[provenance["entity_type"] == entity_type]
    if evidence.empty:
        return frame.copy()
    evidence = (
        evidence.groupby("entity_id", as_index=False)
        .agg(
            evidence_providers=("provider", _joined),
            query_fingerprints=("query_fingerprint", _joined),
            local_fingerprints=("local_fingerprint", _joined),
            payload_hashes=("payload_hash", _joined),
        )
        .rename(columns={"entity_id": id_column})
    )
    return frame.merge(evidence, on=id_column, how="left")


def export_frames(frames: AnalysisFrames) -> dict[str, pd.DataFrame]:
    """Build the four declared Stage 10 tables from the shared frames."""

    return {
        "events": public_frame(_with_evidence(frames.event_snapshots, frames.provenance, "event_snapshot", "event_item_id")),
        "impacts": public_frame(_with_evidence(frames.impacts, frames.provenance, "impact_observation", "impact_item_id")),
        "correlations": public_frame(frames.correlations),
        "quality": public_frame(frames.quality),
    }


def _csv_bytes(frame: pd.DataFrame) -> bytes:
    return frame.to_csv(index=False, lineterminator="\n").encode("utf-8-sig")


def _write_sheet(workbook: Workbook, name: str, rows: list[list[Any]], *, table: bool = False, wrap_last: bool = False) -> None:
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    for row in rows:
        sheet.append([_public_value(value) for value in row])
    if not rows:
        return
    for cell in sheet[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="center")
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
        if wrap_last and row:
            cell = row[-1]
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            sheet.row_dimensions[cell.row].height = min(120, max(15, 15 * math.ceil(len(str(cell.value or "")) / 42)))
    if table and len(rows) > 1:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = max((len(str(cell.value or "")) for cell in column), default=8)
        sheet.column_dimensions[column[0].column_letter].width = min(max(width + 2, 10), 42)


def _xlsx_bytes(result: QueryResult, frames: AnalysisFrames, tables: Mapping[str, pd.DataFrame]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    counts = {
        "event_families": len(frames.event_families),
        "event_snapshots": len(tables["events"]),
        "hazard_snapshots": len(frames.hazards),
        "impact_observations": len(tables["impacts"]),
        "correlation_evidence": len(tables["correlations"]),
    }
    metadata_rows = [["Field", "Value"]]
    metadata_rows.extend([[f"query.{key}", value] for key, value in result.query.as_dict().items()])
    metadata_rows.extend(
        [
            ["query.fingerprint", result.query.fingerprint],
            ["retrieval.provider", result.metadata.provider],
            ["retrieval.retrieved_at", result.metadata.retrieved_at],
            ["retrieval.complete", result.metadata.complete],
            ["retrieval.warnings", "; ".join(public_text(value) for value in result.metadata.warnings)],
            *[[f"count.{key}", value] for key, value in counts.items()],
        ]
    )
    _write_sheet(workbook, "Query Metadata", metadata_rows, wrap_last=True)
    for name, key in (("Events", "events"), ("Impact Observations", "impacts"), ("Correlation Evidence", "correlations"), ("Quality Summary", "quality")):
        frame = tables[key]
        _write_sheet(workbook, name, [list(frame.columns), *frame.values.tolist()], table=True)
    definitions = [
        ["Term", "Definition"],
        ["Event family", "One exact PDC source event family (monty:src_event_id)."],
        ["Event snapshot", "One event STAC item; repeated snapshots are retained as evidence."],
        ["Impact observation", "One original impact item and detail position; original value and unit are preserved."],
        ["PDC exposure", EXPOSURE_NOTE],
        ["Map point", MAP_NOTE],
        ["Present zero", "The source explicitly reported numeric zero; this is not missing data."],
        ["Missing", "No matching observation was retrieved for the requested category."],
        ["Conflicting", "More than one incompatible source value remains visible."],
        ["Latest rule", frames.snapshot_rule],
    ]
    _write_sheet(workbook, "Definitions", definitions, wrap_last=True)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _feature(geometry: Mapping[str, Any], properties: Mapping[str, Any]) -> dict[str, Any]:
    return {"type": "Feature", "geometry": geometry, "properties": {key: _public_value(value) for key, value in properties.items()}}


def _geojson(
    frames: AnalysisFrames,
    *,
    selected_category: str,
    selected_family_key: str | None,
    footprint: Mapping[str, Any] | None,
) -> dict[str, Any]:
    layers = build_map_layers(frames, selected_category=selected_category, selected_family_key=selected_family_key)
    family_context = (
        frames.event_families.set_index("family_key").to_dict("index")
        if not frames.event_families.empty
        else {}
    )
    value_context = (
        layers.event_points.set_index("family_key").to_dict("index")
        if not layers.event_points.empty
        else {}
    )
    features: list[dict[str, Any]] = []
    for layer_name, frame, id_column in (
        ("event_family", layers.event_points, "event_item_id"),
        ("hazard", layers.hazard_points, "hazard_item_ids"),
        ("impact", layers.impact_points, "impact_item_ids"),
    ):
        for row in frame.to_dict("records"):
            if row.get("longitude") is None or row.get("latitude") is None:
                continue
            family = family_context.get(row.get("family_key"), {})
            values = value_context.get(row.get("family_key"), {})
            properties = {
                "layer": layer_name,
                "family_key": row.get("family_key"),
                "source_event_id": row.get("source_event_id") or family.get("source_event_id"),
                "item_ids": row.get(id_column),
                "provider": row.get("provider"),
                "hazard": row.get("hazard_label") or row.get("hazard_labels") or family.get("hazard_labels"),
                "selected_category": selected_category,
                "selected_value": values.get("numeric_value"),
                "selected_unit": values.get("original_unit"),
                "selected_family": row.get("family_key") == selected_family_key if selected_family_key else False,
                "geometry_meaning": MAP_NOTE,
            }
            features.append(_feature({"type": "Point", "coordinates": [row["longitude"], row["latitude"]]}, properties))
    if footprint:
        values = footprint.get("features") if footprint.get("type") == "FeatureCollection" else [footprint]
        for value in values if isinstance(values, list) else []:
            if not isinstance(value, Mapping) or not isinstance(value.get("geometry"), Mapping):
                continue
            properties = dict(value.get("properties") or {})
            properties.update({"layer": "validated_footprint", "family_key": selected_family_key, "geometry_meaning": "Validated PDC footprint asset."})
            features.append(_feature(value["geometry"], properties))
    return {"type": "FeatureCollection", "features": features}


def _provider_fingerprints(result: QueryResult) -> list[dict[str, Any]]:
    metadata = result.provider_metadata or (result.metadata,)
    values = {
        (
            item.provider,
            item.query_fingerprint,
            item.local_fingerprint,
            item.retrieved_at,
        )
        for item in metadata
    }
    return [
        {
            "provider": provider,
            "query_fingerprint": query_fingerprint,
            "local_fingerprint": local_fingerprint,
            "retrieved_at": retrieved_at,
        }
        for provider, query_fingerprint, local_fingerprint, retrieved_at in sorted(
            values,
            key=lambda parts: tuple("" if value is None else str(value) for value in parts),
        )
    ]


def _report(result: QueryResult, frames: AnalysisFrames, counts: Mapping[str, int], selected_category: str, selected_family_key: str | None) -> str:
    query = result.query
    warnings = "\n".join(f"- {public_text(value)}" for value in result.metadata.warnings) or "- None"
    category = frames.category_status[frames.category_status["impact_category"] == selected_category] if not frames.category_status.empty else pd.DataFrame()
    category_summary = (
        "\n".join(f"- {status}: {count:,}" for status, count in category["status"].value_counts().sort_index().items())
        or "- No matching category rows"
    )
    selected = frames.event_families[frames.event_families["family_key"] == selected_family_key] if selected_family_key and not frames.event_families.empty else pd.DataFrame()
    if selected.empty:
        selected_summary = "- None"
    else:
        row = selected.iloc[0]
        selected_summary = "\n".join(
            (
                f"- Title: {row.get('event_title') or 'Untitled'}",
                f"- Source event ID: {row.get('source_event_id') or 'Unavailable'}",
                f"- Hazard: {', '.join(row.get('hazard_labels') or ()) or 'Unspecified'}",
                f"- Countries: {', '.join(row.get('country_codes') or ()) or 'Unspecified'}",
                f"- Snapshots: {int(row.get('snapshot_count') or 0):,}",
            )
        )
    provenance_summary = (
        "\n".join(
            f"- {provider}: {count:,} evidence rows"
            for provider, count in frames.provenance["provider"].value_counts().sort_index().items()
        )
        if not frames.provenance.empty
        else "- No record-level provenance rows"
    )
    fingerprint_summary = "\n".join(
        f"- {item['provider']}: query `{item['query_fingerprint'] or 'n/a'}`; local `{item['local_fingerprint'] or 'n/a'}`"
        for item in _provider_fingerprints(result)
    ) or "- None"
    temporal = (
        temporal_summary(frames, selected_family_key, selected_category).message
        if selected_family_key
        else "No event family was selected for temporal analysis."
    )
    return f"""# PDC evidence report

Generated: {datetime.now(timezone.utc).isoformat()}  
Provider: {result.metadata.provider}  
Retrieval complete: {result.metadata.complete}  
Query fingerprint: `{query.fingerprint}`

## Filters

- Analysis: {query.analysis_mode}
- Country: {query.country_code or "All countries (event/hazard overview only)"}
- Period: {query.period_label}
- Hazard codes: {", ".join(query.hazard_codes) or "All hazards within the bounded period"}
- Impact types: {", ".join(query.impact_types)}
- Categories: {", ".join(query.categories)}
- Source mode: {query.source_mode}
- Display category: {selected_category}
- Selected event family: {selected_family_key or "None"}

## Evidence counts

{chr(10).join(f"- {key.replace('_', ' ').title()}: {value:,}" for key, value in counts.items())}

## Selected category summary

Category: {selected_category}

{category_summary}

## Selected event

{selected_summary}

## Provenance

{provenance_summary}

Provider fingerprints:

{fingerprint_summary}

## Temporal evolution

{temporal}

- {TEMPORAL_NOTE}
- A separate Montandon ingestion timestamp is not available in the retained STAC evidence.

## Warnings

{warnings}

## Interpretation

- {EXPOSURE_NOTE}
- {MAP_NOTE}
- Declared latest-observation rule: {frames.snapshot_rule}.
- Zero, missing, unavailable, not-retrieved, and conflicting values remain distinct.

## Included files

- `events.csv`
- `impact_observations.csv`
- `correlation_evidence.csv`
- `quality_summary.csv`
- `pdc_evidence.xlsx`
- `pdc_evidence.geojson`
- `manifest.json`
- `report.md`
- `report.html`
"""


def build_export_files(
    result: QueryResult,
    *,
    selected_category: str = "people",
    selected_family_key: str | None = None,
    footprint: Mapping[str, Any] | None = None,
) -> dict[str, bytes]:
    """Create all Stage 10 files in memory without exposing paths or credentials."""

    frames = build_analysis_frames(result)
    tables = export_frames(frames)
    counts = {
        "event_families": len(frames.event_families),
        "event_snapshots": len(tables["events"]),
        "hazard_snapshots": len(frames.hazards),
        "impact_observations": len(tables["impacts"]),
        "correlation_evidence": len(tables["correlations"]),
    }
    files = {name: _csv_bytes(tables[key]) for name, key in CSV_NAMES.items()}
    files["pdc_evidence.xlsx"] = _xlsx_bytes(result, frames, tables)
    files["pdc_evidence.geojson"] = json.dumps(
        _geojson(frames, selected_category=selected_category, selected_family_key=selected_family_key, footprint=footprint),
        ensure_ascii=False,
        indent=2,
    ).encode("utf-8")
    report = _report(result, frames, counts, selected_category, selected_family_key)
    charts = report_charts(event_summary(frames), result.query, selected_category)
    for name, (_, svg) in charts.items():
        files[name] = svg.encode("utf-8")
    chart_files = "".join(f"- `{name}`\n" for name in charts)
    chart_md = "".join(f"### {caption}\n\n![{caption}]({name})\n\n" for name, (caption, _) in charts.items())
    files["report.md"] = (report + chart_files + (f"\n## Charts\n\n{chart_md}" if charts else "")).encode("utf-8")
    chart_html = "".join(f"<figure>{svg}<figcaption>{escape(caption)}</figcaption></figure>" for caption, svg in charts.values())
    files["report.html"] = (
        "<!doctype html><meta charset='utf-8'><title>PDC evidence report</title>"
        "<style>body{font:16px Arial,sans-serif;max-width:960px;margin:2rem auto;line-height:1.5}pre{white-space:pre-wrap}"
        "figure{margin:1.5rem 0}figure svg{max-width:100%;height:auto;border:1px solid #E3E6EB;border-radius:8px}"
        "figcaption{color:#5B6472;font-size:14px}</style>"
        f"<pre>{escape(report + chart_files)}</pre>" + (f"<h2>Charts</h2>{chart_html}" if charts else "")
    ).encode("utf-8")
    manifest = {
        "schema_version": 1,
        "software_version": _software_version(),
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "query": result.query.as_dict(),
        "query_fingerprint": result.query.fingerprint,
        "provider_fingerprints": _provider_fingerprints(result),
        "selection": {"category": selected_category, "family_key": selected_family_key},
        "retrieval": {
            "provider": result.metadata.provider,
            "retrieved_at": result.metadata.retrieved_at,
            "complete": result.metadata.complete,
            "pages": result.metadata.pages,
            "returned_count": result.metadata.returned_count,
            "unique_count": result.metadata.unique_count,
            "server_count": result.metadata.server_count,
            "failed_partitions": list(result.metadata.failed_partitions),
            "warnings": [public_text(value) for value in result.metadata.warnings],
        },
        "counts": counts,
        "files": {name: {"bytes": len(data), "sha256": sha256(data).hexdigest()} for name, data in sorted(files.items())},
        "notes": [EXPOSURE_NOTE, MAP_NOTE, TEMPORAL_NOTE, f"Latest observation rule: {frames.snapshot_rule}."],
    }
    files["manifest.json"] = json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True).encode("utf-8")
    return files


def write_export_bundle(result: QueryResult, output_dir: Path | str, **kwargs: Any) -> dict[str, Path]:
    """Write one new export directory; refuse to overwrite existing files."""

    directory = Path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    files = build_export_files(result, **kwargs)
    existing = [directory / name for name in files if (directory / name).exists()]
    if existing:
        raise FileExistsError(f"export files already exist in {directory}")
    written = {}
    for name, data in files.items():
        path = directory / name
        path.write_bytes(data)
        written[name] = path
    return written


def validate_export_bundle(output_dir: Path | str) -> dict[str, int]:
    """Reopen every Stage 10 output and reconcile counts and hashes."""

    directory = Path(output_dir)
    manifest = json.loads((directory / "manifest.json").read_text(encoding="utf-8"))
    expected = manifest["counts"]
    rows = {}
    for name, key in CSV_NAMES.items():
        with (directory / name).open("r", encoding="utf-8-sig", newline="") as handle:
            rows[key] = max(sum(1 for _ in reader(handle)) - 1, 0)
    if rows["events"] != expected["event_snapshots"] or rows["impacts"] != expected["impact_observations"] or rows["correlations"] != expected["correlation_evidence"]:
        raise ValueError("CSV counts do not reconcile with manifest counts")
    geojson = json.loads((directory / "pdc_evidence.geojson").read_text(encoding="utf-8"))
    if geojson.get("type") != "FeatureCollection" or not isinstance(geojson.get("features"), list):
        raise ValueError("GeoJSON output is not a FeatureCollection")
    workbook = load_workbook(directory / "pdc_evidence.xlsx", read_only=True, data_only=False)
    if tuple(workbook.sheetnames) != SHEET_NAMES:
        raise ValueError("Excel sheet contract is incomplete")
    if workbook["Events"].max_row - 1 != rows["events"] or workbook["Impact Observations"].max_row - 1 != rows["impacts"]:
        raise ValueError("Excel counts do not reconcile with CSV counts")
    workbook_text = "\n".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    workbook.close()
    for name, evidence in manifest["files"].items():
        data = (directory / name).read_bytes()
        if len(data) != evidence["bytes"] or sha256(data).hexdigest() != evidence["sha256"]:
            raise ValueError(f"export integrity check failed for {name}")
    for name in ("report.md", "report.html"):
        if not (directory / name).read_text(encoding="utf-8").strip():
            raise ValueError(f"{name} is empty")
    scan = (workbook_text + "\n" + "\n".join(
        (directory / name).read_text(encoding="utf-8", errors="ignore")
        for name in (*CSV_NAMES, "pdc_evidence.geojson", "manifest.json", "report.md", "report.html")
    )).lower()
    if (
        any(value in scan for value in ("authorization: bearer", "montandon_api_token", "pdc_local_export_path"))
        or PATH_RE.search(scan)
        or re.search(r"(?i)\bbearer\s+(?!<redacted>)[^\s]+", scan)
    ):
        raise ValueError("an export contains a credential marker or local path")
    return {**rows, "geojson_features": len(geojson["features"])}
