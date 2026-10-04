"""Focused Stage 9–10 checks; no profiler or full local scan."""

from __future__ import annotations

from dataclasses import replace
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
from urllib.error import HTTPError

from openpyxl import load_workbook

from dashboard.streamlit_app import queries_from_values, query_from_values
from guard_pdc.api import ApiError, PdcApiProvider
from guard_pdc.config import MontandonConfig
from guard_pdc.exports import SHEET_NAMES, build_export_files, public_text, validate_export_bundle, write_export_bundle
from guard_pdc.models import QuerySpec, ValidationError
from guard_pdc.notebook_ui import NotebookExplorer
from guard_pdc.service import PdcEvidenceService, country_group_counts
from tests.test_query_and_api import SequenceOpener
from tests.test_stage6_to_8 import fixture_result


class ExportTests(unittest.TestCase):
    def test_every_output_reopens_and_counts_reconcile(self) -> None:
        result = fixture_result()
        family_key = result.event_families[0].family_key
        with TemporaryDirectory() as temporary:
            paths = write_export_bundle(result, Path(temporary), selected_family_key=family_key)
            self.assertEqual(
                set(paths),
                {
                    "events.csv",
                    "impact_observations.csv",
                    "correlation_evidence.csv",
                    "quality_summary.csv",
                    "pdc_evidence.xlsx",
                    "pdc_evidence.geojson",
                    "manifest.json",
                    "report.md",
                    "report.html",
                    "chart_events_by_month.svg",
                    "chart_events_by_hazard.svg",
                    "chart_top_events.svg",
                },
            )
            counts = validate_export_bundle(temporary)
            self.assertEqual(counts["events"], 1)
            self.assertEqual(counts["impacts"], 2)
            self.assertGreaterEqual(counts["geojson_features"], 3)

            geojson = json.loads(paths["pdc_evidence.geojson"].read_text(encoding="utf-8"))
            for key in ("source_event_id", "hazard", "selected_category", "selected_value", "selected_unit", "provider"):
                self.assertIn(key, geojson["features"][0]["properties"])
            manifest = json.loads(paths["manifest.json"].read_text(encoding="utf-8"))
            self.assertTrue(manifest["provider_fingerprints"])
            report = paths["report.md"].read_text(encoding="utf-8")
            self.assertIn("## Selected category summary", report)
            self.assertIn("## Selected event", report)
            self.assertIn("## Provenance", report)
            # Every chart is linked from the Markdown report and drawn inline in the HTML report.
            html_report = paths["report.html"].read_text(encoding="utf-8")
            for name in [name for name in paths if name.endswith(".svg")]:
                self.assertIn(f"]({name})", report)
                self.assertTrue(paths[name].read_text(encoding="utf-8").startswith("<svg"))
            self.assertEqual(html_report.count("<svg"), 3)

            workbook = load_workbook(paths["pdc_evidence.xlsx"], read_only=False)
            self.assertEqual(tuple(workbook.sheetnames), SHEET_NAMES)
            self.assertEqual(workbook["Events"].freeze_panes, "A2")
            self.assertFalse(workbook["Events"].sheet_view.showGridLines)
            self.assertEqual(workbook["Events"]["A1"].fill.fgColor.rgb, "001F4E78")
            workbook.close()

    def test_exports_omit_credentials_paths_and_raw_pointers(self) -> None:
        result = fixture_result()
        result = replace(
            result,
            metadata=replace(
                result.metadata,
                warnings=(
                    r"source C:\Users\arun.gandhi\Downloads\pdc_data\file.jsonl",
                    "https://example.test/map.json?token=secret-value&part=1",
                ),
            ),
        )
        files = build_export_files(result)
        public = b"\n".join(data for name, data in files.items() if not name.endswith(".xlsx")).decode("utf-8-sig", errors="ignore").lower()
        self.assertNotIn("c:\\users\\", public)
        self.assertNotIn("secret-value", public)
        self.assertNotIn("raw_pointer", public)
        self.assertIn("redacted", public)
        manifest = json.loads(files["manifest.json"])
        self.assertEqual(manifest["counts"]["event_snapshots"], 1)
        self.assertEqual(manifest["counts"]["impact_observations"], 2)

    def test_public_error_text_redacts_paths_and_bearer_tokens(self) -> None:
        message = public_text(r"failed C:\Users\name\pdc.jsonl with Bearer top-secret")
        self.assertNotIn("C:\\Users", message)
        self.assertNotIn("top-secret", message)
        self.assertIn("<local-path>", message)
        self.assertIn("Bearer <redacted>", message)


class DashboardTests(unittest.TestCase):
    def test_country_group_queries_switch_without_retrieval_and_deduplicate_shared_events(self) -> None:
        values = dict(analysis_mode="country_detail", country_code="phl, BGD NPL, phl", year=2024, months=tuple(range(1, 13)), hazard_labels=(), impact_types=("affected_total",), categories=("people",), include_zero_values=True, include_missing_geometry=True)
        queries = queries_from_values(**values)
        self.assertEqual(tuple(q.country_code for q in queries), ("BGD", "NPL", "PHL"))
        self.assertTrue(all(q.source_mode == "api_only" and q.refresh_api_cache for q in queries))
        calls = []
        def fetch(query):
            calls.append(query)
            return replace(fixture_result(), query=query)
        notebook = NotebookExplorer(fetch)
        notebook._render = lambda: None
        notebook.country.value = values["country_code"]
        notebook.specific_months.value = tuple(range(1, 13))
        self.assertEqual(notebook._queries(), queries)
        self.assertIsNotNone(notebook.retrieve_now())
        self.assertEqual(len(calls), 3)
        self.assertEqual(country_group_counts(notebook.results)["event_families"], 1)
        notebook.country_selector.value = "PHL"
        self.assertEqual(notebook.result.query.country_code, "PHL")
        self.assertEqual(len(calls), 3)
        manifest = json.loads(build_export_files(notebook.result)["manifest.json"])
        self.assertEqual(manifest["query"]["country_code"], "PHL")
        previous = notebook.result
        notebook.country.value = "PHL, XX"
        self.assertIsNone(notebook.retrieve_now())
        self.assertIs(notebook.result, previous)
        self.assertEqual(len(calls), 3)
        with self.assertRaises(ValidationError):
            queries_from_values(**{**values, "country_code": "PHL, XX"})
        with self.assertRaises(ValidationError):
            queries_from_values(**{**values, "country_code": ""})
        overview = queries_from_values(**{**values, "analysis_mode": "annual_country_overview", "country_code": ""})
        self.assertIsNone(overview[0].country_code)

    def _offline_app(self):
        """AppTest with year availability and alert-area access stubbed (no network)."""

        from streamlit.testing.v1 import AppTest
        from guard_pdc.maps import AlertAreas

        path = Path(__file__).parents[1] / "dashboard" / "streamlit_app.py"
        patches = (
            patch("guard_pdc.config.MontandonConfig.from_env", return_value=MontandonConfig(api_token="fixture")),
            patch("guard_pdc.api.PdcApiProvider.years_with_events", return_value={2023: True, 2024: True}),
            patch("guard_pdc.maps.fetch_alert_areas", return_value=AlertAreas("missing", "stubbed in tests")),
        )
        return AppTest.from_file(str(path)), patches

    def test_dashboard_country_switch_uses_loaded_group_only(self) -> None:
        calls = []

        def fetch(query, **_options):
            calls.append(query)
            return replace(fixture_result(), query=query)

        app, patches = self._offline_app()
        with patch("guard_pdc.service.retrieve", side_effect=fetch), patches[0], patches[1], patches[2]:
            app.run(timeout=60)
            self.assertEqual(len(app.exception), 0)
            app.multiselect(key="countries").set_value(["PHL", "BGD"]).run(timeout=60)
            slider = next(widget for widget in app.select_slider if widget.label == "Years")
            self.assertEqual(tuple(slider.options), ("2023", "2024"))
            next(button for button in app.button if button.label == "Retrieve data").click().run(timeout=120)
            self.assertEqual(len(app.exception), 0, [item.value for item in app.exception])
            self.assertEqual(len(calls), 2)
            self.assertTrue(all(query.year == 2023 and query.end_year == 2024 for query in calls))
            selector = app.segmented_control(key="view-country")
            self.assertEqual(tuple(selector.options), ("Bangladesh", "Philippines"))  # full names, not ISO3
            selector.set_value("PHL").run(timeout=120)
            self.assertEqual(len(app.exception), 0)
            self.assertEqual(len(calls), 2)

    def test_dashboard_query_is_the_shared_query_contract(self) -> None:
        query = query_from_values(
            analysis_mode="country_detail",
            country_code="phl",
            year=2024,
            months=(2, 1, 1),
            hazard_labels=("Flood",),
            impact_types=("affected_total",),
            categories=("people",),
            include_zero_values=True,
            include_missing_geometry=True,
        )
        expected = QuerySpec(
            analysis_mode="country_detail",
            source_mode="api_only",
            country_code="PHL",
            year=2024,
            months=(1, 2),
            hazard_codes=("MH0600", "FL", "nat-hyd-flo-flo"),
            impact_types=("affected_total",),
            categories=("people",),
            refresh_api_cache=True,
        )
        self.assertEqual(query, expected)
        self.assertEqual(query.fingerprint, expected.fingerprint)

        notebook = NotebookExplorer(lambda _query: fixture_result())
        self.assertEqual(notebook.source_mode.options, [["Live Montandon API", "api_only"]])
        notebook.specific_months.value = (1, 2)
        notebook.hazards.value = ("MH0600",)
        self.assertEqual(notebook._query(), query)

    def test_interactive_queries_use_fresh_api_without_initializing_local_provider(self) -> None:
        notebook = NotebookExplorer(lambda _query: fixture_result())
        notebook.specific_months.value = (1,)
        dashboard_query = query_from_values(
            analysis_mode="country_detail",
            country_code="PHL",
            year=2024,
            months=(1,),
            hazard_labels=(),
            impact_types=("affected_total",),
            categories=("people",),
            include_zero_values=True,
            include_missing_geometry=True,
        )
        page = {"type": "FeatureCollection", "features": [], "links": []}
        for query in (notebook._query(), dashboard_query):
            with self.subTest(interface_query=query), TemporaryDirectory() as temporary:
                root = Path(temporary)
                config = MontandonConfig(
                    api_token="fixture",
                    api_cache_path=root / "api-cache",
                    local_export_path=root / "unused-export.jsonl",
                    local_index_path=root / "unused-index.sqlite",
                )
                opener = SequenceOpener(
                    page, page, page,
                    HTTPError(config.url("search"), 401, "Unauthorized", {}, None),
                )
                provider = PdcApiProvider(config, opener=opener)
                provider.query_events(replace(query, refresh_api_cache=False))
                with patch("guard_pdc.service.PdcLocalProvider", side_effect=AssertionError("Local provider must not be initialized")):
                    service = PdcEvidenceService(config, api_provider=provider)
                    for _ in range(2):
                        result = service.retrieve(query)
                        self.assertEqual(result.metadata.provider, "api")
                    with self.assertRaises(ApiError):
                        service.retrieve(query)
                    self.assertIsNone(service.local_provider)
                self.assertEqual(len(opener.calls), 4)
                self.assertFalse(config.local_index_path.exists())

    def test_dashboard_initial_render_does_not_retrieve(self) -> None:
        app, patches = self._offline_app()
        with patch("guard_pdc.service.retrieve", side_effect=AssertionError("no retrieval on page load")), patches[0], patches[1], patches[2]:
            app.run(timeout=60)
            self.assertEqual(len(app.exception), 0)
            self.assertIn("No retrieval occurs on page load", app.info[0].value)
            self.assertTrue(any("Live Montandon API only" in caption.value for caption in app.sidebar.caption))
            self.assertFalse(any(selectbox.label == "Source" for selectbox in app.selectbox))
            app.session_state["pdc_result"] = fixture_result()
            app.run(timeout=60)
            self.assertEqual(len(app.exception), 0)
            self.assertIn("No retrieval occurs on page load", app.info[0].value)

if __name__ == "__main__":
    unittest.main()
